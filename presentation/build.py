"""Build the editable deck, embedded-font PDF, and standalone HTML preview.

Run from the repository root: python3 presentation/build.py
"""
from pathlib import Path
from io import BytesIO
import base64
import html
import json
import re
import shutil
import subprocess

from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'presentation'
W, H = 960, 540
C = dict(ivory='faf8f4', navy='062344', coral='ff654f', peach='f5bfaa',
         peachlight='fbe5db', blue='b9d6ec', bluelight='e9f1f7', muted='526172', line='dcdedc')
for name, filename in [('regular', 'manrope-regular.ttf'), ('bold', 'manrope-semibold.ttf'), ('extra', 'manrope-extrabold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(ROOT / 'assets' / filename)))
pdfmetrics.registerFont(TTFont('serif', '/usr/share/fonts/truetype/liberation2/LiberationSerif-Italic.ttf'))

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = 'People First AI — ludzka strona transformacji'
prs.core_properties.author = 'People First AI'
pdf = canvas.Canvas(str(OUT / 'People-First-AI.pdf'), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor('People First AI')
pages, slide, elements = [], None, []

def color(value):
    return C.get(value, value)

def rgb(value):
    return RGBColor.from_string(color(value))

def rect(x, y, w, h, fill, circle=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if circle else MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    pdf.setFillColor('#' + color(fill))
    if circle:
        pdf.ellipse(x, H-y-h, x+w, H-y, stroke=0, fill=1)
    else:
        pdf.rect(x, H-y-h, w, h, stroke=0, fill=1)
    elements.append(f'<div style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;background:#{color(fill)};border-radius:{"50%" if circle else "0"}"></div>')

def text(value, x, y, w, size=18, font='regular', fill='navy', leading=1.3):
    # Explicit wrapping makes PDF and PowerPoint share the same line breaks.
    lines = []
    for paragraph in value.split('\n'):
        line = ''
        for word in paragraph.split(' '):
            candidate = (line + ' ' + word).strip()
            if line and pdfmetrics.stringWidth(candidate, font, size) > w - 3:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    h = len(lines)*size*leading + 8
    assert y+h <= H+1, (value, y, h)
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Georgia' if font == 'serif' else 'Manrope'
        p.font.size = Pt(size)
        p.font.bold = font in ('bold', 'extra')
        p.font.italic = font == 'serif'
        p.font.color.rgb = rgb(fill)
        p.line_spacing = Pt(size*leading)
        p.space_before = p.space_after = Pt(0)
        pdf.setFont(font, size)
        pdf.setFillColor('#'+color(fill))
        pdf.drawString(x, H-y-size-i*size*leading, line)
    family = 'Georgia,serif' if font == 'serif' else 'Manrope,sans-serif'
    elements.append(f'<div style="position:absolute;left:{x}px;top:{y}px;width:{w}px;font-size:{size}px;line-height:{leading};font-family:{family};font-weight:{600 if font in ("bold","extra") else 400};font-style:{"italic" if font == "serif" else "normal"};color:#{color(fill)}">'+ '<br>'.join(html.escape(s) for s in lines) + '</div>')

def image(path, x, y, w, h, crop=None):
    im = Image.open(ROOT / path).convert('RGBA')
    if crop:
        im = im.crop(crop)
    ratio = max(w/im.width, h/im.height)
    tw, th = round(w/ratio), round(h/ratio)
    # Top alignment preserves faces in the original hero photo.
    im = im.crop(((im.width-tw)//2, 0, (im.width+tw)//2, th))
    buf = BytesIO()
    im.save(buf, format='PNG')
    slide.shapes.add_picture(BytesIO(buf.getvalue()), Pt(x), Pt(y), Pt(w), Pt(h))
    pdf.drawImage(ImageReader(BytesIO(buf.getvalue())), x, H-y-h, w, h, mask='auto')
    uri = 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode()
    elements.append(f'<img alt="{html.escape(path)}" src="{uri}" style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px">')

def start(kicker, bg='ivory'):
    global slide, elements
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    pdf.setFillColor('#'+color(bg))
    pdf.rect(0, 0, W, H, stroke=0, fill=1)
    elements = [f'<div class="background" style="background:#{color(bg)}"></div>']
    image('assets/logo.png', 48, 26, 178, 28.07)
    text(kicker, 48, 84, 850, 10, 'extra')
    rect(48, 501, 864, .6, 'line')
    text('Technology changes. People matter.', 48, 511, 650, 9, fill='muted')
    text(f'{len(prs.slides):02d} / 16', 863, 510, 70, 10, fill='muted')

def end(note=''):
    slide.notes_slide.notes_text_frame.text = note or 'Treść na podstawie strony głównej People First AI (index.html i app.js).'
    pages.append('<section class="slide">' + ''.join(elements) + '</section>')
    pdf.showPage()

def heading(a, b=None, size=38):
    text(a, 48, 119, 860, size, 'bold', leading=1.12)
    if b:
        text(b, 48, 119+size*1.18, 860, size, 'serif', leading=1.12)

start('KURS O LUDZKIEJ STRONIE AI')
text('W firmie', 48, 126, 420, 39, 'bold')
text('AI First', 48, 174, 420, 43, 'serif')
text('zrób miejsce dla', 48, 238, 420, 36, 'bold')
text('People First.', 48, 284, 420, 43, 'bold')
rect(48, 341, 306, 3, 'coral')
text('Naucz się przeprowadzać ludzi przez zmianę AI.', 48, 376, 388, 20)
image('assets/duet.webp', 492, 122, 420, 365)
rect(506, 444, 151, 30, 'ivory')
text('Ola / liderka', 516, 453, 134, 10, 'bold')
rect(667, 444, 231, 30, 'ivory')
text('Grzegorz / ambasador', 677, 453, 210, 10, 'bold')
end('Otwarcie: główne hasło strony. Fotografia z tej samej sesji co na stronie; stworzona z pomocą AI na podstawie zdjęć prowadzących.')

start('TECHNOLOGIA SIĘ ZMIENIA. LUDZIE MAJĄ ZNACZENIE.')
heading('Dwie perspektywy.', 'Jedna ludzka zmiana.')
text('Z perspektywy lidera i osoby, która doświadcza jej na co dzień.', 48, 237, 850, 18, fill='muted')
for i, (n, label) in enumerate([('5','tygodni wspólnej nauki'),('10','warsztatów na żywo'),('20','godzin praktyki'),('1','Twój plan transformacji')]):
    x = 48+i*220
    text(n, x, 315, 200, 74, 'bold')
    text(label, x, 421, 185, 15)
end()

start('ZMIANA MA DWIE STRONY')
heading('Nowe narzędzia.', 'Te same ludzkie potrzeby.')
for x, fill, person, q, role in [(48,'peachlight','OLA','Jak wesprzeć\nmój zespół?','Perspektywa managerska'),(492,'bluelight','GRZEGORZ','Jak odnaleźć\nwłasne miejsce?','Perspektywa pracownika i ambasadora')]:
    rect(x, 244, 420, 190, fill)
    text(person, x+24, 261, 370, 10, 'extra')
    text(q, x+24, 290, 370, 29, 'serif', leading=1.15)
    text(role, x+24, 395, 370, 12)
text('Wspólna odpowiedź zaczyna się od ludzi.', 48, 456, 864, 17, 'bold')
end()

start('NASZ WSPÓLNY PUNKT WYJŚCIA', 'bluelight')
heading('Najpierw człowiek.', 'Potem framework.')
pillars=[('H','Hear','Wysłuchaj','Rozpoznaj emocje, obawy i potrzeby.'),('U','Understand','Zrozum','Zdiagnozuj gotowość, kulturę i kompetencje.'),('M','Mobilize','Zaangażuj','Włącz liderów i ambasadorów. Daj ludziom wpływ.'),('A','Adapt','Wspieraj adaptację','Rozwijaj kompetencje i bezpieczne eksperymenty.'),('N','Nurture','Rozwijaj','Mierz skutki. Ucz się z sukcesów i niepowodzeń.')]
for i,(letter,en,pl,body) in enumerate(pillars):
    x=48+i*176
    text(letter,x,236,160,68,'extra', 'coral' if i==0 else 'navy')
    text(en,x,333,162,12,'bold')
    text(pl,x,359,162,14,'bold')
    text(body,x,402,151,12)
text('Wysłuchaj. Zrozum. Wróć do ludzi.',48,472,850,12,'bold')
end()

start('DLA TYCH, KTÓRZY PRACUJĄ Z LUDŹMI')
text('Nie musisz być\nekspertem od AI.',48,133,460,38,'bold',leading=1.15)
text('Zacznij od ludzi.',48,235,460,38,'serif')
text('Ten kurs jest dla Ciebie, jeśli pomagasz innym odnaleźć się w zmianie — niezależnie od stanowiska.',48,329,382,19)
roles=['HR i People & Culture','Learning & Development','Menedżerowie i liderzy zespołów','Change managerowie i liderzy transformacji','Ambasadorzy AI i społeczności praktyków']
for i,role in enumerate(roles):
    y=140+i*66
    text(role,509,y,393,17,'bold')
    rect(509,y+52,403,.6,'line')
end()

# Extract the exact course data without running browser-dependent code.
js=(ROOT/'app.js').read_text()
raw=subprocess.check_output(['node','-e',js.split('const programList')[0]+'console.log(JSON.stringify(weeks))'],text=True)
weeks=json.loads(raw)
start('OD ROZUMIENIA DO DZIAŁANIA')
heading('Pięć tygodni.', 'Konkretna zmiana.')
for i,week in enumerate(weeks):
    y=238+i*47
    text(f'0{i+1}',48,y,60,23,'bold',fill='coral')
    text(week['title'],120,y+3,775,20,'bold')
    rect(120,y+37,792,.6,'line')
end()

for i,week in enumerate(weeks):
    start(f'TYDZIEŃ {i+1} / {week["pillar"]}')
    text(week['title'],48,121,855,34,'bold',leading=1.13)
    for j,(title,body,result) in enumerate(week['workshops']):
        x=48+j*444
        rect(x,223,420,252,'peachlight' if j==0 else 'bluelight')
        text(f'WARSZTAT {i*2+j+1:02d}',x+22,241,374,10,'extra')
        text(title,x+22,270,374,22,'bold',leading=1.16)
        text(body,x+22,337,374,14,leading=1.45)
        text('WYCHODZISZ Z',x+22,422,374,9,'extra')
        text(result,x+22,443,374,12,'bold')
    end()

start('POZNAJMY SIĘ / PERSPEKTYWA MANAGERSKA')
image('assets/duet.webp',48,122,343,365,crop=(40,0,820,1024))
text('Aleksandra\nBrzozowska',433,123,475,38,'bold',leading=1.13)
text('Przywództwo. Zmiana. Rozwój ludzi.',433,227,475,17,'serif')
text('Patrzy na zmianę z perspektywy liderów i całej organizacji — z człowiekiem w centrum.',433,282,457,19,'bold')
text('Head of Learning & Development w Ringier Axel Springer Polska, założycielka Future Craft Lab i członkini zarządu Fundacji Digital Poland. Od niemal 20 lat łączy świat technologii, biznesu i ludzi.',433,373,457,14,leading=1.45)
end('Biogram zgodny ze stroną główną. Fotografia kampanii stworzona z pomocą AI na podstawie zdjęć prowadzących.')

start('POZNAJMY SIĘ / PERSPEKTYWA PRACOWNIKA I AMBASADORA')
image('assets/duet.webp',569,122,343,365,crop=(750,0,1536,1024))
text('Grzegorz\nŁukasik',48,123,485,38,'bold',leading=1.13)
text('Praktyka. Społeczność. Codzienna praca.',48,227,482,17,'serif')
text('Patrzy na zmianę z perspektywy pracownika i ambasadora, który pomaga innym oswoić technologię.',48,282,463,19,'bold')
text('Corporate Capabilities Specialist w Ringier Axel Springer Polska, wcześniej Prompt Engineer i developer. Łączy doświadczenie w projektowaniu rozwiązań AI i automatyzacji z edukacją oraz rozwojem kompetencji.',48,373,463,14,leading=1.45)
end('Biogram zgodny ze stroną główną. Fotografia kampanii stworzona z pomocą AI na podstawie zdjęć prowadzących.')

start('WIEDZA, KTÓRA ZOSTAJE W PRAKTYCE','peachlight')
text('Wychodzisz\nz własnym',48,126,465,43,'bold',leading=1.13)
text('planem zmiany.',48,235,470,43,'serif')
text('People-Centred AI Change Plan',48,320,440,17,'bold')
text('Przez cały kurs pracujesz nad jednym projektem — własnym wyzwaniem organizacyjnym lub przygotowanym studium przypadku.',48,368,400,16,leading=1.45)
for i,s in enumerate(['Diagnoza potrzeb, obaw i gotowości','Komunikacja i wsparcie menedżerów','Kompetencje, ambasadorzy i eksperymenty','Pomiar efektów i scenariusze dalszych działań']):
    y=149+i*78
    text(f'0{i+1}',510,y,45,16,'extra',fill='coral')
    text(s,561,y,332,19,'bold')
end()

start('JAK PRACUJEMY')
heading('Razem, na żywo.', 'Przez działanie. Z pomocą AI.')
for i,(title,body) in enumerate([('Razem, na żywo','Kohorta, rozmowy i wymiana doświadczeń.'),('Przez działanie','Ćwiczenia, symulacje i wzajemny feedback.'),('Z pomocą AI','Scenariusze rozmów i warianty komunikacji.')]):
    x=48+i*296
    rect(x,257,272,164,'bluelight' if i==1 else 'peachlight')
    text(f'0{i+1}',x+20,277,230,13,'extra')
    text(title,x+20,312,232,21,'bold')
    text(body,x+20,357,228,15)
text('5 tygodni  /  10 warsztatów na żywo  /  20 godzin praktyki',48,454,855,16,'bold')
end()

start('PEOPLE FIRST. TAKŻE W ŚWIECIE AI.')
rect(0,127,18,334,'peach')
rect(942,127,18,334,'blue')
heading('Zmiana technologiczna.', 'Ludzkie podejście.',size=45)
text('Zrób miejsce dla ludzi w transformacji AI.',48,269,850,23,'bold')
text('Pięć tygodni, dwie perspektywy i Twój plan działania.',48,319,850,19)
rect(48,384,367,42,'bluelight')
rect(63,401,7,7,'coral',circle=True)
text('Przygotowujemy pierwszą edycję',83,396,320,14,'bold')
text('Termin, cena i harmonogram zostaną ogłoszone.',48,445,850,14,fill='muted')
text('Sesja wizerunkowa stworzona z pomocą AI na podstawie zdjęć prowadzących.',48,475,850,8,fill='muted')
end('Nie dodawano niepotwierdzonych terminów, ceny ani adresu zapisów. Strona główna zawiera obecnie formularz podglądowy.')

assert len(prs.slides)==16
prs.save(OUT/'People-First-AI.pptx')
pdf.save()
fonts=[]
for weight,filename in [(400,'manrope-regular.ttf'),(600,'manrope-semibold.ttf'),(800,'manrope-extrabold.ttf')]:
    data=base64.b64encode((ROOT/'assets'/filename).read_bytes()).decode()
    fonts.append(f'@font-face{{font-family:Manrope;font-weight:{weight};src:url(data:font/ttf;base64,{data})}}')
css=''.join(fonts)+'''*{box-sizing:border-box}body{margin:0;background:#202a38;font-family:Manrope,sans-serif}.slide{position:relative;width:960px;height:540px;margin:32px auto;overflow:hidden;box-shadow:0 8px 48px #0004}.background{position:absolute;inset:0}.controls{position:sticky;top:0;z-index:2;background:#062344;color:#faf8f4;padding:12px 24px;text-align:center;font-size:13px}button{font:inherit;padding:6px 14px;cursor:pointer;margin-left:16px}@page{size:960px 540px;margin:0}@media print{body{background:none}.controls{display:none}.slide{margin:0;break-after:page;box-shadow:none}}'''
document='<!doctype html><html lang="pl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>People First AI — prezentacja</title><style>'+css+'</style><body><div class="controls">People First AI · 16 slajdów · ← → nawigacja <button onclick="document.documentElement.requestFullscreen()">Pełny ekran</button><button onclick="print()">Drukuj / PDF</button></div>'+''.join(pages)+'''<script>let i=0;const slides=[...document.querySelectorAll('.slide')];document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown','ArrowLeft','PageUp','Home','End'].includes(e.key)){e.preventDefault();i=e.key==='Home'?0:e.key==='End'?slides.length-1:Math.max(0,Math.min(slides.length-1,i+(['ArrowRight','PageDown'].includes(e.key)?1:-1)));slides[i].scrollIntoView({block:'center'});}});</script></body></html>'''
(OUT/'People-First-AI.html').write_text(document)
fontdir=OUT/'fonts'
fontdir.mkdir(exist_ok=True)
for filename in ['manrope-regular.ttf','manrope-semibold.ttf','manrope-extrabold.ttf','Manrope-OFL.txt']:
    shutil.copy2(ROOT/'assets'/filename,fontdir/filename)
print('Created 16 slides: PowerPoint, PDF, standalone HTML.')
